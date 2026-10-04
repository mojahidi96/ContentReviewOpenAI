from io import BytesIO

from docx import Document
from openpyxl import Workbook

from app.document_loaders.docx_loader import load_docx
from app.document_loaders.excel_loader import load_xlsx
from app.document_loaders.dispatch import validate_signature
from app.document_loaders.base import ExtractionError, ExtractedDocument
from app.services.chunking import split_documents


def test_docx_extracts_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("Quarterly results")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Revenue"
    table.cell(0, 1).text = "$10"
    buffer = BytesIO()
    doc.save(buffer)
    extracted = load_docx(buffer.getvalue())
    assert "Quarterly results" in extracted[0].text
    assert "Revenue | $10" in extracted[0].text


def test_xlsx_rows_retain_sheet_header_and_range():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Risks"
    sheet.append(["Risk", "Impact"])
    sheet.append(["Supply", "High"])
    buffer = BytesIO()
    workbook.save(buffer)
    extracted = load_xlsx(buffer.getvalue())
    assert extracted[0].metadata == {"sheet_name": "Risks", "cell_range": "A2:B2"}
    assert "Risk: Supply" in extracted[0].text


def test_signature_and_chunk_metadata():
    try:
        validate_signature(".pdf", b"not a pdf")
    except ExtractionError as exc:
        assert exc.code == "UNSUPPORTED_FILE"
    chunks = split_documents([ExtractedDocument("one two three", {"page_number": 4})], 8, 2)
    assert chunks[0].metadata["page_number"] == 4
