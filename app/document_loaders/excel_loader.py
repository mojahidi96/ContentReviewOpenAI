from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.document_loaders.base import ExtractedDocument, ExtractionError


def load_xlsx(data: bytes) -> list[ExtractedDocument]:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        chunks: list[ExtractedDocument] = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(value).strip() if value is not None else "" for value in rows[0]]
            for row_number, row in enumerate(rows[1:] if len(rows) > 1 else rows, start=2 if len(rows) > 1 else 1):
                values = []
                for column_number, value in enumerate(row):
                    if value is None:
                        continue
                    label = headers[column_number] if column_number < len(headers) and headers[column_number] else get_column_letter(column_number + 1)
                    values.append(f"{label}: {str(value)[:2000]}")
                if not values:
                    continue
                end_cell = f"{get_column_letter(max(len(row), 1))}{row_number}"
                start_cell = f"A{row_number}"
                chunks.append(
                    ExtractedDocument(
                        text=f"Worksheet: {sheet.title}; row {row_number}\n" + " | ".join(values),
                        metadata={"sheet_name": sheet.title, "cell_range": f"{start_cell}:{end_cell}"},
                    )
                )
        workbook.close()
        if not chunks:
            raise ExtractionError("EMPTY_DOCUMENT", "XLSX contains no extractable cell values.")
        return chunks
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("CORRUPT_DOCUMENT", "XLSX could not be read.") from exc
