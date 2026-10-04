from io import BytesIO

from docx import Document

from app.document_loaders.base import ExtractedDocument, ExtractionError


def load_docx(data: bytes) -> list[ExtractedDocument]:
    try:
        document = Document(BytesIO(data))
        blocks: list[tuple[int, int, str]] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                blocks.append((paragraph._p.getparent().index(paragraph._p), 0, paragraph.text.strip()))
        for table in document.tables:
            position = table._tbl.getparent().index(table._tbl)
            for row_number, row in enumerate(table.rows, start=1):
                values = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
                if any(values):
                    blocks.append((position, row_number, " | ".join(values)))
        blocks.sort(key=lambda block: (block[0], block[1]))
        text = "\n".join(block[2] for block in blocks).strip()
        if not text:
            raise ExtractionError("EMPTY_DOCUMENT", "DOCX contains no extractable text.")
        return [ExtractedDocument(text=text)]
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("CORRUPT_DOCUMENT", "DOCX could not be read.") from exc
