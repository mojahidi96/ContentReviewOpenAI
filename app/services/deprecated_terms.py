import logging
import re
from dataclasses import dataclass
from pathlib import Path
from threading import Lock

from openpyxl import load_workbook

logger = logging.getLogger(__name__)

_DEPRECATED_HEADERS = {"deprecated term", "deprecated", "deprecated terms"}
_REPLACEMENT_HEADERS = {"replacement term", "replacement", "replacement terms"}


@dataclass(frozen=True)
class TermMatch:
    start: int
    end: int
    text: str
    term: str
    replacement: str


def _match_case(matched: str, replacement: str) -> str:
    letters = [c for c in matched if c.isalpha()]
    if len(letters) > 1 and all(c.isupper() for c in letters):
        return replacement.upper()
    if matched[:1].isupper():
        return replacement[:1].upper() + replacement[1:]
    return replacement[:1].lower() + replacement[1:] if replacement[:1].isupper() and not _is_acronym(replacement) else replacement


def _is_acronym(text: str) -> bool:
    return len(text) > 1 and text[:2].isupper()


def read_term_pairs(path: Path) -> list[tuple[str, str]]:
    """Read (deprecated, replacement) pairs from the first sheet; columns are found by header, else A and B."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []
    header = [str(v).strip().lower() if v is not None else "" for v in rows[0]]
    old_col = next((i for i, h in enumerate(header) if h in _DEPRECATED_HEADERS), 0)
    new_col = next((i for i, h in enumerate(header) if h in _REPLACEMENT_HEADERS), 1)
    has_header = any(h in _DEPRECATED_HEADERS | _REPLACEMENT_HEADERS for h in header)
    pairs: list[tuple[str, str]] = []
    for row in rows[1:] if has_header else rows:
        old = str(row[old_col]).strip() if len(row) > old_col and row[old_col] is not None else ""
        new = str(row[new_col]).strip() if len(row) > new_col and row[new_col] is not None else ""
        if old and new and old.lower() != new.lower():
            pairs.append((old, new))
    return pairs


class DeprecatedTerms:
    """Deprecated-term dictionary backed by an Excel file; reloads when the file changes."""

    def __init__(self, path: Path) -> None:
        self._path = path
        self._lock = Lock()
        self._mtime: float | None = None
        self._pattern: re.Pattern[str] | None = None
        self._replacements: dict[str, tuple[str, str]] = {}

    @staticmethod
    def _key(text: str) -> str:
        return " ".join(text.lower().split())

    def _refresh(self) -> None:
        try:
            mtime = self._path.stat().st_mtime
        except OSError:
            if self._mtime is not None or self._pattern is not None:
                logger.warning("Deprecated terms file %s is not readable; check disabled", self._path)
            self._mtime, self._pattern, self._replacements = None, None, {}
            return
        if mtime == self._mtime:
            return
        try:
            pairs = read_term_pairs(self._path)
        except Exception:
            logger.exception("Could not read deprecated terms from %s; keeping previous list", self._path)
            self._mtime = mtime
            return
        replacements = {self._key(old): (old, new) for old, new in pairs}
        # Longest terms first so "man hours" wins over "man".
        ordered = sorted(replacements.values(), key=lambda pair: -len(pair[0]))
        self._pattern = (
            re.compile(
                r"(?<!\w)(?:"
                + "|".join(r"\s+".join(re.escape(w) for w in old.split()) for old, _ in ordered)
                + r")(s?)(?!\w)",
                re.IGNORECASE,
            )
            if ordered
            else None
        )
        self._replacements = replacements
        self._mtime = mtime
        logger.info("Loaded %d deprecated terms from %s", len(replacements), self._path)

    def find(self, content: str) -> list[TermMatch]:
        with self._lock:
            self._refresh()
            pattern, replacements = self._pattern, self._replacements
        if pattern is None:
            return []
        matches: list[TermMatch] = []
        for found in pattern.finditer(content):
            text = found.group(0)
            plural = found.group(1)
            base = self._key(text[: len(text) - len(plural)] if plural else text)
            entry = replacements.get(base)
            if entry is None:  # the term itself ends in "s" (e.g. "Guys")
                base, plural = self._key(text), ""
                entry = replacements.get(base)
            if entry is None:
                continue
            term, replacement = entry
            matches.append(
                TermMatch(found.start(), found.end(), text, term, _match_case(text, replacement) + plural)
            )
        return matches
